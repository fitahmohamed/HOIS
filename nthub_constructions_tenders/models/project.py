# -*- coding: utf-8 -*-
import base64
import io
from datetime import date
from odoo import models, fields, api, _
from odoo.exceptions import UserError

try:
    import openpyxl
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:
    openpyxl = None


class Project(models.Model):
    _inherit = 'project.project'

    offer_count = fields.Integer(string=_('Offers'), compute='get_offer_count')
    num_of_units = fields.Integer(string=_('Number of Units'), required=True)
    is_tender = fields.Boolean(string=_("Tender"))
    approved_tender = fields.Boolean(string=_("Approved Tender"))
    tender_ids = fields.One2many('project.tender', 'project_id', string=_("Tender"), copy=False, ondelete='cascade',
                                 index=True, )
    total_tender_cost = fields.Float(string=_("Total Tender Cost"), store=True, compute='_compute_tender_total')
    total_extra_cost = fields.Float(string=_("Total Extra Cost"), store=True, compute='_compute_tender_total')
    total_cost = fields.Float(string=_("Total Cost"), store=True, compute='_compute_tender_total')
    tender_submission_date = fields.Date(string=_("Tender Submission Date"))
    notes_tender = fields.Text(string=_("Notes"))
    tender_responsible_id = fields.Many2one('hr.employee', string=_("Tender Responsible"),
                                            default=lambda self: self.env.user.employee_id)
    tender_creation_date = fields.Date(string=_("Tender Creation Date"), default=fields.Date.today())
    tender_approval_date = fields.Date(string=_('Tender Approval Date'))
    contract_count = fields.Integer(string=_('Contract'), compute='get_contract_count')
    offer_cost_total = fields.Float(string=_('Offer Cost Total'),
                                    compute='_compute_offer_totals', store=True)
    offer_sell_total = fields.Float(string=_('Total Selling Price'),
                                    compute='_compute_offer_totals', store=True)
    analytic_account_id = fields.Many2one(
        'account.analytic.account',
        string=_('Project Cost Center'),
        copy=False,
    )
    project_overhead_ids = fields.One2many(
        'project.overhead.line', 'project_id', string=_('Overhead Items'),
        help='بنود الأوفر هيد الفعلية على مستوى المشروع (عمالة، مهندسين، '
             'مصاريف، معدات...)، يتم توزيع إجماليها على بنود العرض.')
    total_overhead_cost = fields.Float(string=_("Total Overhead Cost"), store=True, compute='_compute_tender_total')
    project_duration_months = fields.Float(
        string=_('Project Duration (Months)'), default=0.0,
        help='مدة المشروع بالشهور — تُستخدم في حساب بنود الأوفر هيد: الكمية × الشهور × سعر الوحدة'
    )


    @api.onchange('num_of_units')
    def _onchange_num_of_units(self):
        """
        Set the number of units to at least 1
        """
        if self.num_of_units < 1:
            # Set the value to at least 1
            self.num_of_units = 1

    @api.depends('partner_id')
    def _compute_offer_totals(self):
        for rec in self:
            offer = self.env['tender.offer'].search([
                ('project_id', '=', rec.id), ('state', '=', 'confirm'),
            ], limit=1)
            if not offer:
                offer = self.env['tender.offer'].search([
                    ('project_id', '=', rec.id)
                ], order='id desc', limit=1)
            rec.offer_cost_total = offer.total_lines_cost if offer else 0.0
            rec.offer_sell_total = offer.total_sell      if offer else 0.0

    def action_download_tender_template(self):
        """Generate and download a clean Excel import template."""
        if openpyxl is None:
            raise UserError(_('openpyxl is not installed. Contact your administrator.'))

        import io as _io, base64 as _b64
        from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.datavalidation import DataValidation

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Tender Items"

        H     = "1F4E79"
        thin  = Side(style="thin",   color="BDC3C7")
        med   = Side(style="medium", color=H)
        b_all = Border(left=thin, right=thin, top=thin, bottom=thin)
        b_hdr = Border(left=med,  right=med,  top=med,  bottom=med)

        cols = [
            ("كود البند / Code",          14, True),
            ("اسم البند / Name",          30, True),
            ("الوصف / Description",                        34, False),
            ("الكمية / Qty",                           12, True),
            ("وحدة القياس / UoM", 16, False),
            ("النوع / Type",                                16, False),
        ]
        for ci, (hdr, width, req) in enumerate(cols, 1):
            cell = ws.cell(row=1, column=ci, value=hdr)
            cell.font      = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            cell.fill      = PatternFill("solid", fgColor="1A5276" if req else "2980B9")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border    = b_hdr
            ws.column_dimensions[get_column_letter(ci)].width = width
        ws.row_dimensions[1].height = 30

        for row in range(2, 202):
            alt = (row % 2 == 0)
            for ci, (_, _, req) in enumerate(cols, 1):
                cell = ws.cell(row=row, column=ci, value="")
                bg = ("EBF5FB" if alt else "FFFFFF") if req else ("F4F6F7" if alt else "FDFEFE")
                cell.fill      = PatternFill("solid", fgColor=bg)
                cell.font      = Font(name="Calibri", size=10)
                cell.alignment = Alignment(horizontal="left", vertical="center")
                cell.border    = Border(left=thin, right=thin, top=thin, bottom=thin)
            ws.row_dimensions[row].height = 18

        dv = DataValidation(
            type="list",
            formula1='"\u0645\u0648\u0627\u062f,\u0639\u0645\u0627\u0644\u0629,\u0645\u0635\u0631\u0648\u0641\u0627\u062a,\u0645\u0639\u062f\u0627\u062a,\u0645\u0642\u0627\u0648\u0644"',
            allow_blank=True, showDropDown=False,
        )
        ws.add_data_validation(dv)
        dv.sqref = "F2:F201"
        ws.freeze_panes = "A2"
        ws.sheet_view.rightToLeft = True
        ws.sheet_view.rightToLeft = True

        buf = _io.BytesIO()
        wb.save(buf)
        file_b64 = _b64.b64encode(buf.getvalue()).decode()

        att = self.env["ir.attachment"].create({
            "name": "tender_items_template.xlsx",
            "type": "binary",
            "datas": file_b64,
            "mimetype": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "res_model": self._name,
            "res_id": self.id,
        })
        return {
            "type": "ir.actions.act_url",
            "url": f"/web/content/{att.id}?download=true",
            "target": "self",
        }

    def _find_last_breakdown_for_item(self, code, item):
        """آخر بريك داون معمول لنفس البند — بالكود فقط.

        لو البند ملوش كود أو الكود ملوش بريك داون سابق → يعتبر جديد
        ويتعمله بريك داون جديد على طول من غير سؤال.
        """
        if not code:
            return self.env['tender.job.cost']
        return self.env['tender.job.cost'].search(
            [('state', '!=', 'cancel'), ('tender_item_code', '=', code)],
            order='id desc', limit=1)

    def _copy_breakdown_for_line(self, source, line):
        """ينسخ آخر بريك داون بكل تفاصيله (الخامات/العمالة/المعدات/المصروفات/
        المقاولين/الخصومات/الإضافات) وينشئه كبريك داون جديد مربوط بالبند."""
        new_template = self.env['tender.job.cost'].create({
            'name':               line.item_id.name,
            'description':        line.description or line.item_id.name,
            'notes_job':          source.notes_job,
            'company_id':         self.company_id.id,
            'currency_id':        self.currency_id.id,
            'tender_item_name':   line.item_id.name,
            'tender_item_desc':   line.description or '',
            'tender_item_qty':    line.qty,
            'tender_item_uom_id': line.uom_id.id if line.uom_id else False,
            'tender_item_code':   line.reference or line.item_id.code or '',
            'project_id':         self.id,
            'employee_id':        source.employee_id.id if source.employee_id else False,
            'department_id':      source.department_id.id if source.department_id else False,
        })
        # نسخ اللاينات يدوياً (معظم حقول اللاين copy=False فمينفعش نستخدم copy())
        all_lines = self.env['tender.job.cost.line'].search(
            [('job_cost_id', '=', source.id)])
        line_vals = []
        for src in all_lines:
            line_vals.append({
                'job_cost_id':  new_template.id,
                'product_id':   src.product_id.id,
                'description':  src.description,
                'date':         fields.Date.today(),
                'qty':          src.qty,
                # fallback: lines saved without UoM take the product's UoM
                'uom_id':       src.uom_id.id or src.product_id.uom_id.id or False,
                'unit_price':   src.unit_price,
                'cost_actual':  src.cost_actual,
                'cost_per_day': src.cost_per_day,
                'no_of_days':   src.no_of_days,
                'flag':         src.flag,
            })
        if line_vals:
            self.env['tender.job.cost.line'].create(line_vals)
        # ربط البريك داون المنسوخ بسطر التندر — من غير السطر ده الكوبي
        # كان بيتعمل بس التمبلت في التندر بتفضل فاضية
        line.template_id = new_template.id
        return new_template

    def _create_empty_breakdown_for_line(self, line, code):
        """إنشاء بريك داون جديد فاضي للبند وربطه."""
        template = self.env['tender.job.cost'].create({
            'name':               line.item_id.name,
            'description':        line.description or line.item_id.name,
            'company_id':         self.company_id.id,
            'currency_id':        self.currency_id.id,
            'tender_item_name':   line.item_id.name,
            'tender_item_desc':   line.description or '',
            'tender_item_qty':    line.qty,
            'tender_item_uom_id': line.uom_id.id if line.uom_id else False,
            'tender_item_code':   code,
            'project_id':         self.id,
        })
        line.template_id = template.id
        return template

    def action_create_job_cost_templates(self):
        """
        من شاشة مشروع المناقصة — لكل بند ملوش قالب:
        - لو البند ملوش بريك داون سابق → ينشئ له بريك داون جديد فاضي على طول.
        - لو البند له بريك داون سابق (بالكود) → يفتح ويزارد يخليك صاحب القرار:
          خد كوبي من آخر واحد أو أنشئ جديد.
        """
        self.ensure_one()
        if not self.tender_ids:
            raise UserError(_('No items found in the tender.'))

        lines_todo = self.tender_ids.filtered(lambda l: not l.template_id)
        if not lines_todo:
            raise UserError(_('All items already have job cost templates.'))

        choice_vals = []
        for line in lines_todo:
            code = line.reference or line.item_id.code or ''
            source = self._find_last_breakdown_for_item(code, line.item_id)
            if source:
                # موجود قبل كده → القرار للمستخدم في الويزارد
                choice_vals.append((0, 0, {
                    'tender_line_id': line.id,
                    'source_id':      source.id,
                    'choice':         'copy',
                }))
            else:
                # مش موجود → جديد على طول من غير لف ودوران
                self._create_empty_breakdown_for_line(line, code)

        if not choice_vals:
            return True  # كله اتعمله جديد — خلصنا

        wizard = self.env['jobcost.copy.choice.wizard'].create({
            'project_id': self.id,
            'line_ids':   choice_vals,
        })
        return {
            'name':      _('Items with Previous Breakdown — Choose'),
            'type':      'ir.actions.act_window',
            'res_model': 'jobcost.copy.choice.wizard',
            'res_id':    wizard.id,
            'view_mode': 'form',
            'target':    'new',
        }

    def action_create_offer(self):
        """
        Create an offer + a DRAFT version from the tender lines.
        The version stays in draft so the user can review before confirming.
        """
        self.ensure_one()
        if not self.partner_id:
            raise UserError(_('Please set the customer first.'))
        self.state = 'tender'

        version_lines = []
        for line in self.tender_ids:
            version_lines.append((0, 0, {
                'item_id':    line.item_id.id,
                'name':       line.description or line.item_id.name,
                'uom_id':     line.uom_id.id,
                'template_id':line.template_id.id if line.template_id else False,
                'qty':        line.qty,
                'unit_price': line.unit_price,
                'flag':       line.flag,
            }))

        offer = self.env['tender.offer'].create({
            'customer_id':    self.partner_id.id,
            'project_id':     self.id,
            'date':           fields.Date.today(),
        })

        # جلب الإصدار المعتمد أو الأخير عشان نجيب السعر منه
        version = self.env['tender.offer.version'].search([
            ('project_id', '=', self.id),
            ('state', '=', 'confirm'),
        ], order='id desc', limit=1)
        if not version:
            version = self.env['tender.offer.version'].search([
                ('project_id', '=', self.id),
            ], order='id desc', limit=1)

        # بناء lookup: item_id → unit_price من الإصدار
        version_price_map = {}
        if version:
            for vl in version.version_line_ids:
                version_price_map[vl.item_id.id] = vl.unit_price

        for line in self.tender_ids:
            cost = version_price_map.get(line.item_id.id, line.unit_price or 0.0)
            self.env['tender.offer.line'].create({
                'tender_offer_id':    offer.id,
                'item_id':            line.item_id.id,
                'name':               line.description or line.item_id.name,
                'uom_id':             line.uom_id.id,
                'template_id':        line.template_id.id if line.template_id else False,
                'qty':                line.qty,
                'cost_price':         cost,
                'profit_margin_line': 0.0,
                'sell_price':         0.0,
                'flag':               line.flag,
            })
        # إصدار بحالة مسودة — المستخدم يؤكده بنفسه
        self.env['tender.offer.version'].create({
            'customer_id':      self.partner_id.id,
            'project_id':       self.id,
            'state':            'draft',
            'tender_offer_id':  offer.id,
            'date':             fields.Date.today(),
            'version_line_ids': version_lines
        })
        return {
            'name':      _('Tender Offer'),
            'view_mode': 'form',
            'res_model': 'tender.offer',
            'res_id':    offer.id,
            'type':      'ir.actions.act_window',
            'target':    'current',
        }
    def get_offer_count(self):
        """
        Get the number of offers for the project
        """
        offer = self.env['tender.offer'].search([('project_id', '=', self.id)])
        self.offer_count = len(offer)

    def action_open_offer(self):
        """
        Open the offer for the project
        """
        offer = self.env['tender.offer'].search([('project_id', '=', self.id)])
        return {
            'name': _('Offers'),
            'domain': [('project_id', '=', self.id)],
            'view_type': 'form',
            'view_mode': 'form',
            'res_model': 'tender.offer',
            'type': 'ir.actions.act_window',
            'res_id': offer.id,
        }

    def action_confirm_tender(self):
        """This method confirms the tender and updates the state of the project to 'contract'."""
        for project in self:
            if project.offer_count > 0 and project.state != 'contract':
                if not project.partner_id:
                    raise UserError(_('Please select a customer.'))
                for line in project.tender_ids:
                    if line.qty == 0:
                        raise UserError(_('Please enter quantity for %s.') % line.item_id.name)
                if project.date:
                    project.write({'state': 'contract'})
                    # Create main project cost center
                    if not project.analytic_account_id:
                        project_plan = project._ensure_analytic_plan()
                        if project_plan:
                            main_account = self.env['account.analytic.account'].create({
                                'name': project.name,
                                'code': project.name,
                                'partner_id': project.partner_id.id if project.partner_id else False,
                                'plan_id': project_plan.id,
                                'company_id': project.company_id.id or self.env.company.id,
                            })
                            project.analytic_account_id = main_account.id
                    offer = self.env['tender.offer'].search([('project_id', '=', project.id)])
                    offer.ensure_one()
                    if offer.state != 'confirm':
                        for line in offer.order_line_ids:
                            if line.qty == 0:
                                raise UserError(_('Please enter quantity for %s in offer.') % line.item_id.name)
                        offer.write({'state': 'approved'})
                        offer.version_id.write({'state': 'confirm'})
                    project.approved_tender = True
                    project.tender_approval_date = fields.Datetime.now()
                    project.contract_date = fields.Datetime.now()
                    contract_model = self.env['owner.contract']
                    contract_lines = []
                    offer_lines = offer.order_line_ids.filtered(lambda line: line.extra_expenses == False)
                    # extra_expenses_line = offer.order_line_ids.filtered(lambda line: line.extra_expenses == True)
                    for line in offer_lines:
                        analytic_dist = False
                        if line.template_id and line.template_id.analytic_account_id:
                            analytic_dist = {str(line.template_id.analytic_account_id.id): 100.0}
                        line_data = {
                            'item_id':        line.item_id.id,
                            'description':    line.name,
                            'uom_id':         line.uom_id.id,
                            'quantity':       line.qty,
                            'cost_price':     line.cost_price,
                            'profit_margin':  line.profit_margin_line,
                            'sell_price':     line.sell_price,
                            'price_unit':     line.sell_price,
                            'template_id':    line.template_id.id if line.template_id else False,
                            'percentage':     0,
                            'amount':         line.qty * line.sell_price,
                            'analytic_distribution': analytic_dist,
                        }
                        contract_lines.append((0, 0, line_data))
                    contract_vals = {
                        'project_id': project.id,
                        'partner_id': project.partner_id.id,
                        'date': project.date_start,
                        'end_date': project.date,
                        # 'extra_expenses': extra_expenses_line.sub_total,
                        'received_date': date.today(),
                        'currency_id': project.currency_id.id,
                        'down_payment_percentage': offer.down_payment_percentage,
                        'down_payment': offer.down_payment,
                        'is_owner': True,
                        'state': 'draft',
                        'owner_contract_line_ids': contract_lines,
                    }
                    contract_model.create(contract_vals)
                else:
                    raise UserError(_("Please set the end date for the tender."))
            return True

    def action_open_contract(self):
        """Open the contract records related to the current project."""
        contracts = self.env['owner.contract'].search([('project_id', '=', self.id), ('is_owner', '=', True)])
        return {
            'type': 'ir.actions.act_window',
            'name': 'Contract',
            'res_model': 'owner.contract',
            'view_type': 'form',
            'view_mode': 'form',
            'domain': [('id', 'in', contracts.ids)],
            'res_id': contracts.ids[0],
            'target': 'current',
        }


    def get_contract_count(self):
        """Calculate and set the count of contracts associated with the project."""
        for rec in self:
            contracts = rec.env['owner.contract'].search([
                ('project_id', '=', rec.id), ('is_owner', '=', True)])
            rec.contract_count = len(contracts)

    def _ensure_analytic_plan(self):
        """Ensure the project has an analytic plan under a root Projects plan."""
        self.ensure_one()
        Plan = self.env['account.analytic.plan']
        root_plan = Plan.search([
            ('name', 'in', ['المشاريع', 'Projects']),
            ('parent_id', '=', False),
        ], limit=1)
        if not root_plan:
            root_plan = Plan.create({'name': 'المشاريع'})
        project_plan = Plan.search([
            ('name', '=', self.name),
            ('parent_id', '=', root_plan.id),
        ], limit=1)
        if not project_plan:
            project_plan = Plan.create({
                'name': self.name,
                'parent_id': root_plan.id,
            })
        return project_plan


    @api.depends('tender_ids.qty', 'tender_ids.unit_price', 'project_extras_ids.cost',
                 'project_overhead_ids.amount', 'project_duration_months')
    def _compute_tender_total(self):
        """Compute the total cost of the tender"""
        for rec in self:
            rec.total_tender_cost = sum([(p.qty * p.unit_price) for p in rec.tender_ids])
            rec.total_extra_cost = sum([(p.cost) for p in rec.project_extras_ids])
            rec.total_overhead_cost = sum([(p.amount) for p in rec.project_overhead_ids])
            rec.total_cost = rec.total_tender_cost + rec.total_extra_cost

    def _recompute_offers_direct_costs(self):
        """إعادة حساب التكلفة المباشرة لعروض السعر المرتبطة بهذا المشروع عند
        تغيير بنود الأوفر هيد."""
        for project in self:
            offers = self.env['tender.offer'].search([('project_id', '=', project.id)])
            if offers:
                offers._recompute_direct_costs()
